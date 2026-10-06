"""Import and quality-control helpers for experimental campaigns."""
from __future__ import annotations

import csv
import io
import math
from collections import Counter

from openpyxl import load_workbook
import xlsxwriter

from materia.science import REQUIRED, parse_measurements


def parse_measurement_file(raw: bytes, filename: str) -> list[dict]:
    suffix=filename.lower().rsplit('.',1)[-1] if '.' in filename else ''
    if suffix=='csv':
        return parse_measurements(raw)
    if suffix!='xlsx':
        raise ValueError('Formats acceptés : CSV UTF-8 ou classeur Excel .xlsx.')
    if len(raw)>5_000_000:
        raise ValueError('Classeur trop volumineux : maximum 5 Mo.')
    try:
        workbook=load_workbook(io.BytesIO(raw),read_only=True,data_only=True)
        sheet=workbook.active
        iterator=sheet.iter_rows(values_only=True)
        headers=[str(value).strip() if value is not None else '' for value in next(iterator)]
        if not set(REQUIRED)<=set(headers):
            missing=sorted(set(REQUIRED)-set(headers))
            raise ValueError('Colonnes manquantes : '+', '.join(missing)+'.')
        out=io.StringIO(); writer=csv.writer(out); writer.writerow(headers)
        count=0
        for row in iterator:
            if all(value is None or str(value).strip()=='' for value in row):
                continue
            writer.writerow(list(row)); count+=1
            if count>10000:
                raise ValueError('Maximum 10 000 mesures par fichier.')
        return parse_measurements(out.getvalue().encode('utf-8'))
    except ValueError:
        raise
    except Exception as exc:
        raise ValueError('Classeur Excel illisible ou endommagé.') from exc


def quality_report(rows: list[dict]) -> dict:
    if not rows:
        raise ValueError('La campagne ne contient aucune mesure.')
    groups=sorted({r['experiment_id'] for r in rows}); flags=[]; details=[]
    duplicate_keys=[key for key,count in Counter((r['experiment_id'],r['time_days']) for r in rows).items() if count>1]
    if duplicate_keys:
        flags.append({'severity':'bloquant','message':f'{len(duplicate_keys)} temps dupliqué(s) dans une même expérience.'})
    for group in groups:
        part=sorted((r for r in rows if r['experiment_id']==group),key=lambda r:r['time_days'])
        times=[float(r['time_days']) for r in part]; values=[float(r['modulus_MPa']) for r in part]
        has_baseline=any(abs(t)<1e-12 for t in times)
        if len(set(times))<3:
            flags.append({'severity':'bloquant','message':f'{group} : moins de trois temps distincts.'})
        if not has_baseline:
            flags.append({'severity':'attention','message':f'{group} : aucun temps initial t=0.'})
        jumps=[]
        for before,after in zip(values,values[1:]):
            if before>0: jumps.append(abs(after-before)/before*100)
        if jumps and max(jumps)>50:
            flags.append({'severity':'attention','message':f'{group} : variation supérieure à 50 % entre deux temps ; vérifier unités et éprouvettes.'})
        details.append({'experiment_id':group,'points':len(part),'distinct_times':len(set(times)),
                        'start_days':min(times),'end_days':max(times),'has_t0':has_baseline,
                        'module_min_mpa':min(values),'module_max_mpa':max(values),
                        'max_step_change_pct':max(jumps) if jumps else 0})
    def texts(field):
        return sorted({str(r[field]).strip() for r in rows if r.get(field) not in (None,'')})
    def numbers(field):
        return sorted({float(r[field]) for r in rows if r.get(field) not in (None,'')})
    for field,label in (('temperature_C','température d\'exposition'),('humidity_RH','humidité relative'),
                        ('thickness_mm','épaisseur'),('material','matériau'),('protocol','protocole'),
                        ('source','source')):
        if any(r.get(field) in (None,'') for r in rows):
            flags.append({'severity':'bloquant','message':f'Champ {label} absent sur une ou plusieurs mesures ; compléter la campagne.'})
    conditions={
        'materials':texts('material'),
        'temperatures_C':numbers('temperature_C'),
        'humidity_RH':numbers('humidity_RH'),
        'thickness_mm':numbers('thickness_mm'),
        'protocols':texts('protocol'),
        'sources':texts('source'),
    }
    blocking=sum(f['severity']=='bloquant' for f in flags)
    return {'rows':len(rows),'experiments':len(groups),'details':details,'conditions':conditions,
            'flags':flags,'blocking_count':blocking,'ready_for_review':blocking==0,
            'status':'Prêt pour revue scientifique' if blocking==0 else 'Correction requise avant revue'}


def measurement_template_xlsx() -> bytes:
    output=io.BytesIO(); workbook=xlsxwriter.Workbook(output,{'in_memory':True})
    data=workbook.add_worksheet('Mesures'); guide=workbook.add_worksheet('Mode d’emploi')
    header=workbook.add_format({'bold':True,'font_color':'#FFFFFF','bg_color':'#235ABE','text_wrap':True,'align':'center'})
    number=workbook.add_format({'num_format':'0.000'}); note=workbook.add_format({'text_wrap':True,'valign':'top'})
    data.write_row(0,0,REQUIRED,header)
    samples=[]
    for experiment,temp,rh,thickness in [('EXPERIENCE-A',40,60,1.0),('EXPERIENCE-B',60,75,1.5)]:
        for hours in (0,240,480,960):
            e0=2000.; modulus=e0*math.exp(-hours*(0.00005+temp/4_000_000+rh/8_000_000)/thickness)
            samples.append([experiment,hours,'hours',round(modulus,3),'MPa',e0,'MPa',temp,rh,thickness,23,'MATERIAL-ID','tensile','Référence complète ou DOI',f'Tableau/figure, {experiment}, t={hours} h'])
    for row_index,row in enumerate(samples,1):
        data.write_row(row_index,0,row)
    data.add_table(0,0,len(samples),len(REQUIRED)-1,{'name':'MesuresMateria','style':'Table Style Medium 2','columns':[{'header':h} for h in REQUIRED]})
    data.freeze_panes(1,2); data.set_column(0,0,20); data.set_column(1,12,17); data.set_column(13,14,34)
    guide.set_column('A:A',28); guide.set_column('B:B',90); guide.hide_gridlines(2)
    guide.write_row(0,0,['Champ','Règle'],header)
    rules=[('Une ligne','Une mesure moyenne à un temps donné. Utilisez des identifiants d’expérience différents pour les campagnes indépendantes.'),
           ('Unités','Temps : hours, days ou seconds. Module : MPa, GPa ou Pa. Materia normalise ensuite en jours et MPa.'),
           ('Protocole','La version actuelle accepte tensile. Décrivez la norme, la vitesse et la température dans la source et l’emplacement.'),
           ('Répétitions','Conservez les données brutes par éprouvette dans votre archive de laboratoire. Ce fichier importe une valeur par expérience et temps.'),
           ('Validation','L’import reste en attente jusqu’à la revue d’un enseignant ou responsable scientifique.')]
    for index,row in enumerate(rules,1): guide.write_row(index,0,row,note); guide.set_row(index,42)
    workbook.close(); return output.getvalue()

"""Freeze predictions before experiments and score them afterwards without data leakage."""
from __future__ import annotations

from datetime import datetime, timezone
from hashlib import sha256
import csv
import io
import json
import math
import re

import numpy as np
from openpyxl import load_workbook

from materia.specimen_analysis import analyse_specimens

ACCEPTANCE_TARGETS={'minimum_points':5,'maximum_mape_pct':10.,'maximum_absolute_bias_pct':5.}
VALIDATION_COLUMNS=('time_days','modulus_mpa','replicate_id','lot_id','specimen_id')
MAX_VALIDATION_ROWS=10_000


def _canonical(payload: dict) -> bytes:
    return json.dumps(payload,ensure_ascii=False,sort_keys=True,separators=(',',':'),allow_nan=False).encode()


def _detached(value):
    """Return a JSON-safe deep copy so later edits cannot silently alter a snapshot."""
    return json.loads(json.dumps(value,ensure_ascii=False,allow_nan=False))


def _validated_curve(snapshot: dict) -> tuple[np.ndarray,np.ndarray,np.ndarray,np.ndarray]:
    """Validate the complete frozen curve before interpolation or scoring."""
    try:
        times=np.asarray(snapshot['time_years'],dtype=float)*365.25
        central=np.asarray(snapshot['central_mpa'],dtype=float)
        lower=np.asarray(snapshot['lower_mpa'],dtype=float)
        upper=np.asarray(snapshot['upper_mpa'],dtype=float)
    except (KeyError,TypeError,ValueError) as exc:
        raise ValueError('Le fichier figé ne contient pas une courbe numérique complète.') from exc
    size=len(times)
    if size<2 or not (len(central)==len(lower)==len(upper)==size):
        raise ValueError('Les quatre séries de la prédiction figée doivent avoir la même longueur.')
    if not all(np.all(np.isfinite(values)) for values in (times,central,lower,upper)):
        raise ValueError('La prédiction figée contient une valeur non finie.')
    if abs(float(times[0]))>1e-9 or np.any(np.diff(times)<=0):
        raise ValueError('Le temps de la prédiction doit commencer à zéro et augmenter strictement.')
    if np.any(central<=0) or np.any(lower<0) or np.any(upper<=0):
        raise ValueError('Les modules de la prédiction doivent être positifs.')
    if np.any(lower>central) or np.any(central>upper):
        raise ValueError('La courbe centrale doit rester entre ses deux bornes.')
    return times,central,lower,upper


def parse_validation_file(raw: bytes, filename: str) -> list[dict]:
    """Read the small prospective-validation format from CSV or XLSX."""
    if not raw:
        raise ValueError('Le fichier de mesures est vide.')
    if len(raw)>5_000_000:
        raise ValueError('Fichier trop volumineux : maximum 5 Mo.')
    suffix=filename.lower().rsplit('.',1)[-1] if '.' in filename else ''
    rows=[]
    if suffix=='csv':
        try:
            text=raw.decode('utf-8-sig')
        except UnicodeDecodeError as exc:
            raise ValueError('Le CSV doit être encodé en UTF-8.') from exc
        reader=csv.DictReader(io.StringIO(text))
        if not reader.fieldnames or not {'time_days','modulus_mpa'}<=set(reader.fieldnames):
            raise ValueError('Colonnes obligatoires : time_days et modulus_mpa.')
        rows=list(reader)
    elif suffix=='xlsx':
        try:
            workbook=load_workbook(io.BytesIO(raw),read_only=True,data_only=True)
            iterator=workbook.active.iter_rows(values_only=True)
            headers=[str(value).strip() if value is not None else '' for value in next(iterator)]
            if not {'time_days','modulus_mpa'}<=set(headers):
                raise ValueError('Colonnes obligatoires : time_days et modulus_mpa.')
            rows=[dict(zip(headers,row)) for row in iterator if not all(value in (None,'') for value in row)]
        except ValueError:
            raise
        except Exception as exc:
            raise ValueError('Classeur Excel illisible ou endommagé.') from exc
    else:
        raise ValueError('Formats acceptés : CSV UTF-8 ou Excel .xlsx.')
    if not rows:
        raise ValueError('Le fichier ne contient aucune mesure.')
    if len(rows)>MAX_VALIDATION_ROWS:
        raise ValueError(f'Maximum {MAX_VALIDATION_ROWS} mesures par validation.')
    observations=[]
    for index,row in enumerate(rows,2):
        try:
            day=float(row['time_days']); modulus=float(row['modulus_mpa'])
        except (TypeError,ValueError,KeyError) as exc:
            raise ValueError(f'Ligne {index} : temps ou module invalide.') from exc
        if not math.isfinite(day) or not math.isfinite(modulus) or day<0 or modulus<=0:
            raise ValueError(f'Ligne {index} : temps positif ou nul et module strictement positif requis.')
        replicate=str(row.get('replicate_id') or row.get('specimen_id') or f'M{index-1}').strip()[:100]
        specimen=str(row.get('specimen_id') or replicate).strip()[:100]
        lot=str(row.get('lot_id') or '').strip()[:100]
        if not lot:
            match=re.match(r'^(L[^-]+)-',specimen,re.IGNORECASE)
            lot=match.group(1) if match else 'Lot non renseigné'
        observations.append({'time_days':day,'modulus_mpa':modulus,'replicate_id':replicate,
                             'lot_id':lot,'specimen_id':specimen})
    return observations


def freeze_prediction(result: dict, title: str, data_cutoff: str | None = None) -> dict:
    """Create an immutable, self-verifying snapshot before target observations are known."""
    manifest=result.get('manifest') or {}
    if not result.get('time') or not result.get('modulus') or len(result['time'])!=len(result['modulus']):
        raise ValueError('Résultat incomplet : courbe impossible à figer.')
    snapshot={
        'format':'materia-blind-prediction-v1',
        'title':str(title).strip() or 'Prédiction aveugle',
        'frozen_at':datetime.now(timezone.utc).isoformat(),
        'data_cutoff':data_cutoff or datetime.now(timezone.utc).date().isoformat(),
        'model':manifest.get('model'),
        'model_fingerprint':result.get('fingerprint'),
        'evidence_level':manifest.get('evidence_level'),
        'source':_detached(manifest.get('source')),
        'inputs':_detached(manifest.get('inputs')),
        'uncertainty':_detached(manifest.get('uncertainty')),
        'time_years':[float(value) for value in result['time']],
        'central_mpa':[float(value) for value in result['modulus']],
        'lower_mpa':[float(value) for value in result['lower']],
        'upper_mpa':[float(value) for value in result['upper']],
    }
    _validated_curve(snapshot)
    snapshot['curve_domain_days']=[0.,float(snapshot['time_years'][-1])*365.25]
    snapshot['lock_sha256']=sha256(_canonical(snapshot)).hexdigest()
    return snapshot


def verify_snapshot(snapshot: dict) -> bool:
    expected=snapshot.get('lock_sha256')
    if not isinstance(expected,str):
        return False
    payload={key:value for key,value in snapshot.items() if key!='lock_sha256'}
    try:
        valid=sha256(_canonical(payload)).hexdigest()==expected
        if valid:
            _validated_curve(snapshot)
        return valid
    except (ValueError,TypeError,OverflowError):
        return False


def compare_with_experiment(snapshot: dict, observations: list[dict]) -> dict:
    """Compare a frozen curve with later measurements inside the frozen horizon."""
    if not verify_snapshot(snapshot):
        raise ValueError('La prédiction a été modifiée après son gel ou son empreinte est invalide.')
    times,central,lower,upper=_validated_curve(snapshot)
    if not len(observations):
        raise ValueError('Ajoutez au moins une observation expérimentale.')
    validated=[]
    for index,row in enumerate(observations,1):
        day=float(row['time_days']); observed=float(row['modulus_mpa'])
        if not math.isfinite(day) or not math.isfinite(observed) or day<0 or observed<=0:
            raise ValueError('Temps et module expérimentaux invalides.')
        if day>times[-1]:
            raise ValueError('Une observation dépasse l’horizon qui avait été figé.')
        replicate=str(row.get('replicate_id') or row.get('specimen_id') or f'M{index}')
        specimen=str(row.get('specimen_id') or replicate)
        lot=str(row.get('lot_id') or 'Lot non renseigné')
        validated.append({'time_days':day,'modulus_mpa':observed,
                          'replicate_id':replicate,'lot_id':lot,'specimen_id':specimen})
    specimen_analysis=analyse_specimens(validated)
    details=[]
    for day in sorted({row['time_days'] for row in validated}):
        part=[row for row in validated if row['time_days']==day]
        measurements=np.asarray([row['modulus_mpa'] for row in part],dtype=float)
        observed=float(np.mean(measurements)); sd=float(np.std(measurements,ddof=1)) if len(measurements)>1 else 0.
        predicted=float(np.interp(day,times,central)); low=float(np.interp(day,times,lower)); high=float(np.interp(day,times,upper))
        details.append({
            'time_days':day,'observed_mpa':observed,'predicted_mpa':predicted,
            'lower_mpa':low,'upper_mpa':high,'error_mpa':predicted-observed,
            'relative_error_pct':abs(predicted-observed)/observed*100,
            'covered':bool(low<=observed<=high),'replicates':len(part),
            'observed_sd_mpa':sd,
        })
    errors=np.asarray([row['error_mpa'] for row in details],dtype=float)
    observed=np.asarray([row['observed_mpa'] for row in details],dtype=float)
    mape=float(np.mean(np.abs(errors)/observed)*100)
    bias=float(np.mean(errors)); bias_pct=abs(bias)/float(np.mean(observed))*100
    enough=len(details)>=ACCEPTANCE_TARGETS['minimum_points']
    central_pass=enough and mape<=ACCEPTANCE_TARGETS['maximum_mape_pct'] and bias_pct<=ACCEPTANCE_TARGETS['maximum_absolute_bias_pct']
    interval_validated=bool((snapshot.get('uncertainty') or {}).get('predictive_interval_validated'))
    if not enough:
        verdict='POINTS_INSUFFISANTS'
        verdict_label=f"Ajouter des mesures : {len(details)}/{ACCEPTANCE_TARGETS['minimum_points']} points"
    elif not central_pass:
        verdict='COURBE_CENTRALE_REJETEE_SUR_CETTE_CAMPAGNE'
        verdict_label='La courbe centrale dépasse les objectifs d’erreur sur cette campagne'
    elif not interval_validated:
        verdict='COURBE_CENTRALE_ACCEPTABLE_SUR_UNE_CAMPAGNE'
        verdict_label='Courbe centrale acceptable localement ; généralisation et intervalle encore non validés'
    else:
        verdict='VALIDATION_LOCALE_REUSSIE'
        verdict_label='Objectifs d’erreur atteints ; contrôler encore plusieurs campagnes indépendantes'
    return {
        'snapshot_sha256':snapshot['lock_sha256'],
        'count':len(details),
        'observation_count':len(validated),
        'distinct_timepoints':len(details),
        'mae_mpa':float(np.mean(np.abs(errors))),
        'rmse_mpa':float(np.sqrt(np.mean(errors**2))),
        'nrmse_pct':float(np.sqrt(np.mean(errors**2))/np.mean(observed)*100),
        'mape_pct':mape,
        'bias_mpa':bias,
        'absolute_bias_pct':bias_pct,
        'interval_coverage_pct':float(np.mean([row['covered'] for row in details])*100),
        'validation':{
            'verdict':verdict,'label':verdict_label,
            'targets':ACCEPTANCE_TARGETS,
            'enough_points':enough,'central_curve_pass':central_pass,
            'predictive_interval_evaluable':interval_validated,
            'scope':'Les métriques donnent le même poids à chaque temps, même si le nombre de répétitions diffère. Résultat sur cette campagne uniquement ; trois campagnes indépendantes minimum sont requises pour conclure à la généralisation.',
        },
        'details':details,
        'raw_observations':specimen_analysis['observations'],
        'specimen_analysis':specimen_analysis,
        'interpretation':'Comparaison aveugle : la prédiction et ses sources ont été figées avant l’ajout de ces observations.',
    }
